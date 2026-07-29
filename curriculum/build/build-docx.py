#!/usr/bin/env python3
"""
build-docx — annexure.json -> editable Word document.

Reads the SAME document model that the PDF is rendered from (produced by
build-annexure.js), so the two deliverables cannot drift apart. Tables are real
Word tables with repeating header rows, so the file stays editable for whoever
assembles the final proposal.

    python3 build-docx.py
"""
import json
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "annexure.json")
OUT = os.path.join(HERE, "..", "Annexure-A-Syllabus-and-Lab.docx")

# palette — matches the deck / PDF
INK = RGBColor(0x15, 0x1B, 0x26)
MUTED = RGBColor(0x57, 0x62, 0x74)
FAINT = RGBColor(0x87, 0x91, 0xA3)
ACC = "0E7490"
ACC_INK = RGBColor(0x0A, 0x55, 0x67)
GOV = RGBColor(0x3D, 0x5A, 0x80)
BORDER = "D9DFEA"
BORDER_STRONG = "C4CDDB"
SURF2 = "F4F7FB"
SURF3 = "EDF1F7"
ACC_SOFT = "E2F1F4"

SANS = "Segoe UI"
MONO = "Consolas"

CONTENT_MM = 210 - 13 - 13  # A4 width minus L/R margins


# ------------------------------------------------------------------ helpers --
# OOXML validates child ORDER inside pPr / tcPr / tblPr. Appending raw elements
# produces a file Word silently "repairs" (dropping the formatting) or refuses to
# open. Every raw element below is therefore inserted at its schema position.

PPR_ORDER = [
    "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
    "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd",
    "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
]
TCPR_ORDER = [
    "w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders",
    "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
    "w:vAlign", "w:hideMark",
]
TBLPR_ORDER = [
    "w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual",
    "w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc",
    "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout",
    "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription",
]
# edge order is fixed inside a *Borders element
EDGE_ORDER = ["top", "left", "bottom", "right", "insideH", "insideV"]


def get_or_insert(parent, name, order):
    """Return child `name` of `parent`, creating it at its schema position."""
    found = parent.find(qn(name))
    if found is not None:
        return found
    el = OxmlElement(name)
    pos = order.index(name)
    for child in parent:
        for i, n in enumerate(order):
            if child.tag == qn(n):
                if i > pos:
                    child.addprevious(el)
                    return el
                break
    parent.append(el)
    return el


def set_edges(borders, edges):
    """edges: list of (edge, val, sz, color), written in schema edge order."""
    for edge in EDGE_ORDER:
        for e_name, val, sz, color in edges:
            if e_name != edge:
                continue
            el = OxmlElement("w:" + edge)
            el.set(qn("w:val"), val)
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:color"), color)
            borders.append(el)


def shade_el(parent, fill, order):
    shd = get_or_insert(parent, "w:shd", order)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def para_border(p, edge, color, sz=6, space=2):
    pPr = p._p.get_or_add_pPr()
    pBdr = get_or_insert(pPr, "w:pBdr", PPR_ORDER)
    el = OxmlElement("w:" + edge)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pBdr.append(el)


def para_shade(p, fill):
    shade_el(p._p.get_or_add_pPr(), fill, PPR_ORDER)


def cell_shade(cell, fill):
    shade_el(cell._tc.get_or_add_tcPr(), fill, TCPR_ORDER)


def cell_border(cell, edge, color, sz=4, val="single"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = get_or_insert(tcPr, "w:tcBorders", TCPR_ORDER)
    set_edges(borders, [(edge, val, sz, color)])


def table_borders(table):
    tblPr = table._tbl.tblPr
    b = get_or_insert(tblPr, "w:tblBorders", TBLPR_ORDER)
    set_edges(b, [
        ("top", "single", 6, BORDER_STRONG),
        ("left", "none", 0, "auto"),
        ("bottom", "single", 4, BORDER),
        ("right", "none", 0, "auto"),
        ("insideH", "single", 4, BORDER),
        ("insideV", "none", 0, "auto"),
    ])


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trPr.append(h)


def add_field(p, instr):
    r = p.add_run()
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "end")
    r._r.append(f1)
    r._r.append(it)
    r._r.append(f2)
    return r


def put_runs(p, runs, size=None, color=None, bold_color=None):
    """Render an inline run list into a paragraph."""
    for r in runs:
        run = p.add_run(r["s"])
        run.bold = bool(r.get("b"))
        run.italic = bool(r.get("i"))
        if size:
            run.font.size = size
        if r.get("c"):
            run.font.name = MONO
            run.font.color.rgb = ACC_INK
            if size:
                run.font.size = Pt(size.pt - 0.7)
        elif r.get("b") and bold_color is not None:
            run.font.color.rgb = bold_color
        elif r.get("i") and color is None:
            run.font.color.rgb = MUTED
        elif color is not None:
            run.font.color.rgb = color
    return p


def plain(runs):
    return "".join(r["s"] for r in runs)


def col_widths(block, total_mm):
    """Proportional column widths from content length, sublinear so one long
    column doesn't crush the rest."""
    ncol = len(block["header"])
    scores = []
    for i in range(ncol):
        lens = [len(plain(block["header"][i]))]
        for row in block["rows"]:
            lens.append(len(plain(row[i])))
        mx = max(lens)
        mean = sum(lens) / len(lens)
        scores.append(max(1.0, 0.45 * mx + 0.55 * mean) ** 0.72)
    tot = sum(scores)
    frac = [s / tot for s in scores]
    # clamp then renormalise
    lo, hi = 0.06, 0.50
    frac = [min(hi, max(lo, f)) for f in frac]
    tot = sum(frac)
    return [total_mm * f / tot for f in frac]


# --------------------------------------------------------------------- build --

doc = Document()

sec = doc.sections[0]
sec.page_width = Mm(210)
sec.page_height = Mm(297)
sec.top_margin = Mm(15)
sec.bottom_margin = Mm(18)
sec.left_margin = Mm(13)
sec.right_margin = Mm(13)
sec.footer_distance = Mm(9)

normal = doc.styles["Normal"]
normal.font.name = SANS
normal.font.size = Pt(9.5)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.14

for name, size, color, bold in (
    ("Heading 1", 13, ACC_INK, True),
    ("Heading 2", 11, INK, True),
    ("Heading 3", 10, GOV, True),
    ("Heading 4", 9.5, MUTED, True),
):
    st = doc.styles[name]
    st.font.name = SANS
    st.font.size = Pt(size)
    st.font.color.rgb = color
    st.font.bold = bold
    st.paragraph_format.space_before = Pt(12)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.keep_with_next = True

with open(SRC, encoding="utf-8") as f:
    data = json.load(f)

# ---- footer -----------------------------------------------------------------
fp = sec.footer.paragraphs[0]
fp.paragraph_format.tab_stops.add_tab_stop(Mm(CONTENT_MM / 2), WD_TAB_ALIGNMENT.CENTER)
fp.paragraph_format.tab_stops.add_tab_stop(Mm(CONTENT_MM), WD_TAB_ALIGNMENT.RIGHT)
para_border(fp, "top", BORDER, sz=4, space=4)


def foot_run(text, color=FAINT, bold=False):
    r = fp.add_run(text)
    r.font.name = MONO
    r.font.size = Pt(6.5)
    r.font.color.rgb = color
    r.font.bold = bold
    return r


foot_run("BGR Neo · Centre of Excellence", ACC_INK, True)
foot_run("\tAnnexure A — Syllabus, Timeline & Lab Requirement (Rev 2)")
foot_run("\tPage ")
pf = add_field(fp, " PAGE ")
pf.font.name = MONO
pf.font.size = Pt(6.5)
pf.font.color.rgb = FAINT
foot_run(" of ")
nf = add_field(fp, " NUMPAGES ")
nf.font.name = MONO
nf.font.size = Pt(6.5)
nf.font.color.rgb = FAINT

# ---- masthead ---------------------------------------------------------------
mh = doc.add_paragraph()
mh.paragraph_format.space_before = Pt(0)
mh.paragraph_format.space_after = Pt(1)
para_border(mh, "top", ACC, sz=20, space=1)
r = mh.add_run("ANNEXURE A")
r.font.name = MONO
r.font.size = Pt(7.5)
r.font.bold = True
r.font.color.rgb = ACC_INK

t = doc.add_paragraph()
t.paragraph_format.space_before = Pt(2)
t.paragraph_format.space_after = Pt(3)
r = t.add_run("Syllabus, Timeline & Lab Requirement")
r.font.size = Pt(19)
r.font.bold = True
r.font.color.rgb = INK

for line in data.get("meta", []):
    m = doc.add_paragraph()
    m.paragraph_format.space_before = Pt(0)
    m.paragraph_format.space_after = Pt(0)
    r = m.add_run(line)
    r.font.name = MONO
    r.font.size = Pt(7.5)
    r.font.color.rgb = MUTED

spacer = doc.add_paragraph()
spacer.paragraph_format.space_after = Pt(2)
para_border(spacer, "bottom", BORDER, sz=4, space=6)

# ---- body -------------------------------------------------------------------
h1_seen = 0
n_tables = 0

for b in data["blocks"]:
    typ = b["type"]

    if typ in ("h1", "h2", "h3", "h4", "h5", "h6"):
        lvl = min(4, int(typ[1]))
        p = doc.add_paragraph(style="Heading %d" % lvl)
        if lvl == 1:
            h1_seen += 1
            if h1_seen > 1:
                p.paragraph_format.page_break_before = True
            para_border(p, "bottom", ACC, sz=12, space=3)
            for r in b["runs"]:
                run = p.add_run(r["s"].upper())
                run.font.bold = True
        else:
            put_runs(p, b["runs"])
        continue

    if typ == "hr":
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(6)
        para_border(p, "bottom", BORDER, sz=4, space=1)
        continue

    if typ == "p":
        p = doc.add_paragraph()
        for i, line in enumerate(b["lines"]):
            if i:
                p.add_run().add_break()
            put_runs(p, line)
        continue

    if typ == "quote":
        for i, line in enumerate(b["lines"]):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Mm(3)
            p.paragraph_format.right_indent = Mm(2)
            p.paragraph_format.space_before = Pt(6 if i == 0 else 0)
            p.paragraph_format.space_after = Pt(6 if i == len(b["lines"]) - 1 else 2)
            para_border(p, "left", ACC, sz=18, space=6)
            para_shade(p, ACC_SOFT)
            put_runs(p, line)
        continue

    if typ in ("ul", "ol"):
        style = "List Bullet" if typ == "ul" else "List Number"
        for item in b["items"]:
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(2)
            put_runs(p, item)
        continue

    if typ == "table":
        n_tables += 1
        ncol = len(b["header"])
        wide = ncol >= 5
        fs = Pt(8.0 if wide else 8.5)
        hs = Pt(7.4 if wide else 7.9)

        tbl = doc.add_table(rows=1, cols=ncol)
        tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
        tbl.autofit = False   # emits <w:tblLayout w:type="fixed"/> itself
        table_borders(tbl)

        widths = col_widths(b, CONTENT_MM)

        hdr = tbl.rows[0]
        repeat_header(hdr)
        for i, cell in enumerate(hdr.cells):
            cell.width = Mm(widths[i])
            cell_border(cell, "bottom", ACC, sz=8)
            cell_shade(cell, SURF3)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            if b["align"][i] == "right":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif b["align"][i] == "center":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in b["header"][i]:
                run = p.add_run(r["s"])
                run.bold = True
                run.font.size = hs
                run.font.color.rgb = INK

        for ri, row in enumerate(b["rows"]):
            cells = tbl.add_row().cells
            for i, cell in enumerate(cells):
                cell.width = Mm(widths[i])
                if ri % 2 == 1:
                    cell_shade(cell, SURF2)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(1.6)
                p.paragraph_format.space_after = Pt(1.6)
                p.paragraph_format.line_spacing = 1.05
                if b["align"][i] == "right":
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif b["align"][i] == "center":
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                put_runs(p, row[i], size=fs)

        after = doc.add_paragraph()
        after.paragraph_format.space_after = Pt(6)
        continue

doc.core_properties.title = "Annexure A — Syllabus, Timeline & Lab Requirement"
doc.core_properties.subject = "BGR Neo Centre of Excellence — AI for Robotics"
doc.core_properties.author = "BGR Neo"

doc.save(OUT)
print("wrote %s (%.0f KB) · %d tables"
      % (os.path.relpath(OUT), os.path.getsize(OUT) / 1024, n_tables))
